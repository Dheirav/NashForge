"""
The phase-2 paper and report as one Springer LNCS-style Word document.

    venv/bin/python scripts/phase2/make_springer.py

Writes docs/NashForge_LNCS.docx from scripts/phase2/content.py. Figures come
from make_figures.py (run it first); the arena numbers are recomputed from
results/chipzen/matches each time, so re-running after a season round updates
the tables. make_latex.py renders the same content with the llncs class.
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402

from content import AUTHOR, EMAIL, INSTITUTE, TITLE, build  # noqa: E402

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
FIGS = os.path.join(ROOT, "docs", "figures", "phase2")
OUT = os.path.join(ROOT, "docs", "NashForge_LNCS.docx")

doc = Document()
for section in doc.sections:
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin = section.right_margin = Cm(3.0)
    section.top_margin = section.bottom_margin = Cm(2.8)
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.05
figure_no = table_no = 0


def para(text, size=10, bold=False, italic=False, align=None, space_after=4, indent=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.5)
    return p


para(TITLE, size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
para(AUTHOR, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para(f"{INSTITUTE}\n{EMAIL}", size=9, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

for block in build():
    kind = block[0]
    if kind == "abstract":
        para("Abstract.", size=9, bold=True, space_after=2)
        para(block[1], size=9, space_after=4)
    elif kind == "keywords":
        para("Keywords: " + block[1] + ".", size=9, space_after=10)
    elif kind in ("h1", "h2"):
        p = doc.add_paragraph()
        run = p.add_run(block[1])
        run.bold = True
        run.font.size = Pt(12 if kind == "h1" else 10)
        p.paragraph_format.space_before = Pt(10 if kind == "h1" else 6)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
    elif kind == "p":
        para(block[1], indent=True)
    elif kind == "figure":
        _, name, caption, width = block
        figure_no += 1
        path = os.path.join(FIGS, name)
        if os.path.exists(path):
            doc.add_picture(path, width=Cm(width))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        para(f"Fig. {figure_no}. {caption}", size=9, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    elif kind == "table":
        _, caption, header, rows = block
        table_no += 1
        para(f"Table {table_no}. {caption}", size=9, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
        t = doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        for i, h in enumerate(header):
            cell = t.rows[0].cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(h)
            run.bold = True
            run.font.size = Pt(9)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                run = cells[i].paragraphs[0].add_run(str(v))
                run.font.size = Pt(9)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
    elif kind == "refs":
        for i, r in enumerate(block[1], 1):
            p = para(f"{i}. {r}", size=9, space_after=2)
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.first_line_indent = Cm(-0.6)

doc.save(OUT)
words = sum(len(p.text.split()) for p in doc.paragraphs)
print(f"wrote {os.path.relpath(OUT, ROOT)}: {words:,} words, {figure_no} figures, {table_no} tables")
