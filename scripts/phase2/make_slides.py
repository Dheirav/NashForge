"""
The four demo slides the brief asks for, from the repository's figures.

    venv/bin/python scripts/phase2/make_slides.py

Writes docs/NashForge_Phase2.pptx: Title; Objective; Overall architecture;
Module design with snapshots and results. Run make_figures.py first.
"""
import json
import os
import sys

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..")))

from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.util import Emu, Inches, Pt  # noqa: E402

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
FIGS = os.path.join(ROOT, "docs", "figures", "phase2")
OUT = os.path.join(ROOT, "docs", "NashForge_Phase2.pptx")
sys.path.insert(0, os.path.dirname(__file__))
from arena_stats import arena_summary  # noqa: E402

INK = RGBColor(0x2F, 0x3A, 0x44)
ACCENT = RGBColor(0x2F, 0x48, 0x58)
SOFT = RGBColor(0x5C, 0x64, 0x6A)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]


def text(slide, x, y, w, h, lines, size=18, bold=False, color=INK, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines if isinstance(lines, list) else [lines]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = "Calibri"
        if align:
            p.alignment = align
    return box


def title(slide, t, sub=None):
    text(slide, 0.6, 0.35, 12.1, 0.9, t, size=30, bold=True, color=ACCENT)
    if sub:
        text(slide, 0.6, 1.1, 12.1, 0.5, sub, size=14, color=SOFT)
    line = slide.shapes.add_shape(1, Inches(0.6), Inches(1.6), Inches(12.1), Emu(20000))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()


def picture(slide, name, x, y, w=None, h=None):
    path = os.path.join(FIGS, name)
    if not os.path.exists(path):
        return
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    slide.shapes.add_picture(path, Inches(x), Inches(y), **kw)


A = arena_summary()
eq = json.load(open(os.path.join(ROOT, "results", "cfr", "experiments", "eq200_vs_eq40.json")))

# 1. Title
s = prs.slides.add_slide(blank)
text(s, 0.8, 2.0, 11.7, 1.6, "NashForge", size=54, bold=True, color=ACCENT)
text(s, 0.8, 3.2, 11.7, 1.2, "One instrument for three learning families in heads-up no-limit hold'em,\nand what a live arena adds",
     size=24, color=INK)
text(s, 0.8, 4.9, 11.7, 1.4, ["Dheirav Prakash",
                             "Department of Computer Science and Engineering, College of Engineering Guindy, Anna University",
                             "CS23E02 Artificial Intelligence, mini project phase 2, September 2026"], size=16, color=SOFT)

# 2. Objective
s = prs.slides.add_slide(blank)
title(s, "Objective", "The question, the constraint, and what phase 2 adds")
text(s, 0.8, 1.9, 6.0, 5.0, [
    "Compare three ways of learning to play poker on ONE instrument:",
    "   regret minimisation (MCCFR), evolutionary search, deep RL (PPO)",
    "",
    "Same engine, same abstraction, same 40,000-hand evaluation with error bars,",
    "so that a difference between families is a measurement and not an artefact.",
    "",
    "Phase 2:",
    "  - the comparison itself, on a converged panel",
    "  - two abstraction findings (six buckets; estimator noise)",
    "  - a C++ core: 32.5x, a solver in 109 s",
    "  - the solver deployed as a rated bot on a public arena,",
    "    every decision logged, and the weakest point fixed",
], size=16)
text(s, 7.2, 1.9, 5.5, 5.0, [
    "Base papers",
    "Lanctot et al. 2009: Monte Carlo CFR (the solver)",
    "Johanson et al. 2013: evaluating abstractions (the findings)",
    "Ganzfried and Sandholm 2013: action translation (the bridge)",
    "",
    "Contribution beyond them",
    "The instrument; the three-way comparison; the noise",
    "diagnosis of a bucket sweep; a native core; and live",
    "rated play used as an instrument in its own right.",
], size=15, color=INK)

# 3. Architecture
s = prs.slides.add_slide(blank)
title(s, "Overall architecture", "engine and abstraction shared; three families; one internal instrument, two external")
picture(s, "fig_architecture.png", 2.2, 1.75, h=5.5)

# 4. Modules with snapshots and results
s = prs.slides.add_slide(blank)
title(s, "Module design, snapshots and results",
      f"arena: {A['matches']} matches, {A['won']} won, {A['hands']:,} hands, {A['net']:+,} chips; decisions in under a millisecond")
picture(s, "fig_arena_modules.png", 0.4, 1.75, w=6.4)
picture(s, "fig_panel.png", 7.0, 1.75, h=2.6)
picture(s, "snap_decision_log.png", 0.4, 5.0, w=6.4)
picture(s, "fig_arena_depth.png", 7.0, 4.55, h=2.6)
text(s, 11.5, 4.6, 1.8, 2.7, [
    "Solver over PPO +75,",
    "over evolution +212",
    "BB/100",
    "",
    f"200 samples: {eq['mean']:+.1f}",
    "BB/100",
    "",
    "Deep stacks lose,",
    "shallow win:",
    "the re-raise gap",
], size=10, color=INK)

prs.save(OUT)
print("wrote", os.path.relpath(OUT, ROOT))
